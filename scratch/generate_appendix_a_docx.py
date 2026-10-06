import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_appendix_a():
    doc = docx.Document()
    
    # 1. Setup Margins (KU Thesis Standard: Left 1.5", Top 1.5", Right 1.0", Bottom 1.0")
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69) # A4
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    
    # Normal style setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'TH Sarabun New'
    normal_style.font.size = Pt(16)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    # Helper to set run font
    def set_font(run, font_name='TH Sarabun New', size=16, bold=False, italic=False, color=RGBColor(0, 0, 0)):
        run.font.name = font_name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rPr.append(rFonts)

    def add_para(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=16, indent=0, space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if indent > 0:
            p.paragraph_format.first_line_indent = Inches(indent)
        run = p.add_run(text)
        set_font(run, size=size, bold=bold)
        return p

    def add_placeholder_box(fig_no, fig_title, fig_desc):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(5.7)
        
        # Border
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for border_name in ['top', 'left', 'bottom', 'right']:
            edge = OxmlElement(f'w:{border_name}')
            edge.set(qn('w:val'), 'dashed')
            edge.set(qn('w:sz'), '8')
            edge.set(qn('w:space'), '0')
            edge.set(qn('w:color'), 'A0AEC0')
            tcBorders.append(edge)
        tcPr.append(tcBorders)
        
        # Shading
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F8FAFC')
        tcPr.append(shd)
        
        # Content inside box
        p_box = cell.paragraphs[0]
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_box.paragraph_format.space_before = Pt(28)
        p_box.paragraph_format.space_after = Pt(28)
        
        r1 = p_box.add_run("[ พื้นที่สำหรับวางรูปภาพ ]\n")
        set_font(r1, size=15, bold=True, color=RGBColor(71, 85, 105))
        
        r2 = p_box.add_run("(คลิกที่นี่หรือลบข้อความนี้เพื่อแทรกรูปภาพ screenshot)")
        set_font(r2, size=13, italic=True, color=RGBColor(148, 163, 184))
        
        # Figure caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap_label = p_cap.add_run(f"ภาพที่ {fig_no} ")
        set_font(r_cap_label, size=16, bold=True)
        r_cap_title = p_cap.add_run(fig_title)
        set_font(r_cap_title, size=16, bold=False)
        
        # Explanatory note
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(12)
        r_note = p_note.add_run(f"(คำอธิบายรูป: {fig_desc})")
        set_font(r_note, size=14, italic=True, color=RGBColor(100, 116, 139))

    # ==================== PAGE 1: COVER PAGE ====================
    for _ in range(7):
        doc.add_paragraph()
        
    p_cov1 = doc.add_paragraph()
    p_cov1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov1 = p_cov1.add_run("ภาคผนวก ก")
    set_font(r_cov1, size=18, bold=True)
    
    p_cov2 = doc.add_paragraph()
    p_cov2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov2 = p_cov2.add_run("คู่มือการใช้งานระบบแชทบอทวินิจฉัยโรคข้าวและแนะนำวิธีการรักษา\nผ่านแอปพลิเคชันไลน์")
    set_font(r_cov2, size=18, bold=True)
    
    doc.add_page_break()
    
    # ==================== PAGE 2: MANUAL BODY ====================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    r_title = p_title.add_run("คู่มือการใช้งานระบบแชทบอทวินิจฉัยโรคข้าวและแนะนำวิธีการรักษา\nผ่านแอปพลิเคชันไลน์")
    set_font(r_title, size=18, bold=True)
    
    add_para(
        "คู่มือการใช้งานระบบแชทบอทวินิจฉัยโรคข้าวและแนะนำวิธีการรักษา แบ่งออกเป็น 2 ส่วน ได้แก่:",
        indent=0.5, space_after=2
    )
    add_para("1) ส่วนสำหรับเกษตรกรและผู้ใช้งานทั่วไป: การใช้งานแชทบอทผ่าน LINE Official Account (@325wtzzd)", indent=0.75, space_after=2)
    add_para("2) ส่วนสำหรับผู้ดูแลระบบ (Admin): การดูแลระบบเวิร์กโฟลว์ (n8n) และฐานข้อมูลความรู้ (Pinecone)", indent=0.75, space_after=12)
    
    # -------------------- ก.1 --------------------
    add_para("ก.1 ส่วนสำหรับเกษตรกรและผู้ใช้งานทั่วไป", bold=True, size=16, space_before=6, space_after=4)
    
    # ก.1.1
    add_para("ก.1.1 การเพิ่มเพื่อนและเริ่มต้นใช้งานแชทบอท", bold=True, size=16, indent=0.25, space_after=3)
    add_para(
        "ผู้ใช้สามารถเริ่มต้นใช้งานได้โดยการสแกนคิวอาร์โค้ด (QR Code) หรือค้นหาด้วย LINE ID: @325wtzzd เมื่อกดเพิ่มเพื่อน บอทจะส่งข้อความทักทายต้อนรับ และแสดงแถบเมนูด้านล่างของหน้าจอแชท",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-1",
        "หน้าจอเมื่อเพิ่มเพื่อนและเริ่มต้นใช้งานแชทบอท",
        "ภาพหน้าจอแอปพลิเคชัน LINE เมื่อเพิ่มเพื่อนสำเร็จ มีข้อความต้อนรับและแถบเมนูขึ้นที่ด้านล่าง"
    )
    
    # ก.1.2
    add_para("ก.1.2 การใช้งานแถบเมนูด้านล่าง (Rich Menu)", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "แถบเมนูด้านล่างหน้าจอแชท มี 3 ปุ่มให้กดใช้งาน ได้แก่ 1) ถ่ายภาพตรวจโรคข้าว 2) ส่งภาพตรวจโรคข้าว และ 3) วิธีใช้งานแชทบอท เพื่อให้ผู้ใช้กดใช้งานได้อย่างสะดวก",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-2",
        "หน้าจอแถบเมนูด้านล่าง (Rich Menu) ของห้องแชท LINE",
        "ภาพหน้าจอแสดงแถบเมนูด้านล่าง 3 ปุ่ม: ถ่ายภาพตรวจโรคข้าว, ส่งภาพตรวจโรคข้าว และวิธีใช้งานแชทบอท"
    )
    
    # ก.1.3
    add_para("ก.1.3 การส่งรูปภาพใบข้าวเพื่อตรวจโรค", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "ผู้ใช้สามารถกดปุ่ม “ถ่ายภาพตรวจโรคข้าว” เพื่อถ่ายรูปใบข้าวจากกล้อง หรือกด “ส่งภาพตรวจโรคข้าว” เพื่อเลือกรูปจากมือถือ แล้วกดส่งเข้ามาในแชทเพื่อให้บอทตรวจหาโรค",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-3",
        "การเลือกและส่งรูปภาพใบข้าวเข้าห้องแชท LINE",
        "ภาพหน้าจอขณะผู้ใช้เลือกรูปใบข้าวที่มีอาการของโรค แล้วกดส่งเข้ามาในห้องแชท"
    )
    
    # ก.1.4
    add_para("ก.1.4 ผลการตรวจโรคและคำแนะนำการรักษา", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "เมื่อบอทตรวจเสร็จ จะส่งข้อความบอกชื่อโรคข้าวที่พบ เปอร์เซ็นต์ความมั่นใจ และวิธีดูแลรักษาตามคำแนะนำของกรมการข้าวกลับมาให้ทันที",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-4",
        "ผลการตรวจโรคข้าวและคำแนะนำการรักษาที่บอทตอบกลับ",
        "ภาพหน้าจอข้อความผลลัพธ์ที่แชทบอทส่งกลับมา แสดงชื่อโรค เปอร์เซ็นต์ความมั่นใจ และวิธีรักษาโรคข้าว"
    )
    
    # ก.1.5
    add_para("ก.1.5 การพิมพ์ถามข้อมูลโรคข้าวเพิ่มเติม", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "ผู้ใช้สามารถพิมพ์ถามเรื่องโรคข้าวหรือวิธีรักษาเข้ามาในแชทได้โดยตรง เช่น “โรคไหม้ข้าวรักษายังไง” บอทจะค้นหาข้อมูลและตอบคำแนะนำที่ถูกต้องกลับมาให้ทันที",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-5",
        "การพิมพ์ถามเรื่องโรคข้าวและคำตอบจากบอท",
        "ภาพหน้าจอขณะผู้ใช้พิมพ์ถามเกี่ยวกับโรคข้าว และแชทบอทตอบคำแนะนำการรักษากลับมา"
    )
    
    # ก.1.6
    add_para("ก.1.6 การแจ้งเตือนเมื่อรูปภาพไม่ถูกต้องหรือไม่ชัดเจน", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "ถ้าผู้ใช้ส่งรูปที่ไม่ใช่ใบข้าว หรือภาพมองไม่เห็นรอยโรค บอทจะแจ้งเตือนให้ถ่ายรูปใบข้าวที่มีรอยโรคให้ชัดเจนในระยะใกล้ แล้วส่งเข้ามาใหม่อีกครั้ง",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-6",
        "ข้อความแจ้งเตือนเมื่อรูปภาพไม่ถูกต้องหรือไม่ชัดเจน",
        "ภาพหน้าจอแสดงข้อความแจ้งเตือนของบอท เมื่อผู้ใช้ส่งภาพอื่นที่ไม่ใช่ใบข้าว หรือภาพมองไม่เห็นรอยโรค"
    )
    
    # -------------------- ก.2 --------------------
    add_para("ก.2 ส่วนสำหรับผู้ดูแลระบบ (Admin)", bold=True, size=16, space_before=10, space_after=4)
    
    # ก.2.1
    add_para("ก.2.1 การเข้าสู่ระบบจัดการ n8n", bold=True, size=16, indent=0.25, space_after=3)
    add_para(
        "ผู้ดูแลระบบสามารถเปิดเว็บเบราว์เซอร์ เข้าสู่ระบบ n8n โดยกรอกอีเมลและรหัสผ่าน แล้วกดปุ่ม “Sign in” เพื่อเข้าสู่ระบบ",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-7",
        "หน้าจอเข้าสู่ระบบจัดการ n8n",
        "ภาพหน้าจอเข้าสู่ระบบ n8n สำหรับให้ผู้ดูแลระบบกรอกอีเมลและรหัสผ่าน"
    )
    
    # ก.2.2
    add_para("ก.2.2 ผังการทำงานของระบบแชทบอท (Workflow Canvas)", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "หน้าจอแสดงผังการทำงานของระบบบน n8n ซึ่งเชื่อมต่อขั้นตอนการรับภาพจาก LINE, การตรวจหาโรคข้าวด้วยโมเดล AI, การค้นหาข้อมูลการรักษา และการส่งข้อความตอบกลับไปยังผู้ใช้",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-8",
        "ผังการทำงานของระบบแชทบอทบน n8n",
        "ภาพหน้าจอผังการเชื่อมต่อขั้นตอนต่าง ๆ บน n8n แสดงเส้นทางการทำงานตั้งแต่รับข้อมูลจนถึงตอบกลับผู้ใช้"
    )
    
    # ก.2.3
    add_para("ก.2.3 การตรวจสอบประวัติการทำงาน (Executions)", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "ผู้ดูแลระบบสามารถคลิกที่เมนู “Executions” เพื่อดูประวัติการทำงานย้อนหลัง ตรวจสอบสถานะว่าทำงานสำเร็จหรือมีข้อผิดพลาด และคลิกดูรายละเอียดของแต่ละขั้นตอนได้",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-9",
        "หน้าจอแสดงประวัติการทำงาน (Executions) บน n8n",
        "ภาพหน้าจอแสดงประวัติการทำงานย้อนหลังในหน้า Executions และการคลิกดูรายละเอียดในแต่ละขั้นตอน"
    )
    
    # ก.2.4
    add_para("ก.2.4 การจัดการฐานข้อมูลความรู้บน Pinecone", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "ผู้ดูแลระบบสามารถเข้าเว็บ Pinecone เพื่อตรวจสอบสถานะของฐานข้อมูล “rice-knowledge” ดูจำนวนข้อมูลที่จัดเก็บ และดูรายการเนื้อหาโรคข้าวแต่ละรายการได้",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-10",
        "หน้าจอจัดการฐานข้อมูล rice-knowledge บน Pinecone",
        "ภาพหน้าจอบนเว็บ Pinecone แสดงสถานะฐานข้อมูล จำนวนข้อมูล และรายการเนื้อหาความรู้โรคข้าว"
    )
    
    # ก.2.5
    add_para("ก.2.5 การอัปเดตข้อมูลความรู้โรคข้าวเข้าสู่ฐานข้อมูล", bold=True, size=16, indent=0.25, space_before=6, space_after=3)
    add_para(
        "หากต้องการเพิ่มหรืออัปเดตข้อมูลโรคข้าว ให้วางไฟล์เอกสารไว้ในโฟลเดอร์ data/knowledge/ แล้วเปิด Terminal รันคำสั่ง python scripts/build_rag_index.py ระบบจะนำเข้าข้อมูลชุดใหม่เข้าสู่ฐานข้อมูลโดยอัตโนมัติ",
        indent=0.5, space_after=8
    )
    add_placeholder_box(
        "ก-11",
        "หน้าต่างการรันคำสั่งอัปเดตข้อมูลความรู้โรคข้าวเข้าสู่ฐานข้อมูล",
        "ภาพหน้าจอ Terminal ขณะรันคำสั่งอัปเดตข้อมูล แสดงขั้นตอนการอ่านไฟล์และการนำเข้าข้อมูลขึ้นระบบสำเร็จ"
    )
    
    # Output path
    out_docx = "Project/Appendix_A.docx"
    try:
        doc.save(out_docx)
        print(f"Successfully generated {out_docx}!")
    except PermissionError:
        out_docx = "Project/Appendix_A_updated.docx"
        doc.save(out_docx)
        print(f"Appendix_A.docx is open in Word. Saved to {out_docx} instead!")

if __name__ == "__main__":
    create_appendix_a()
